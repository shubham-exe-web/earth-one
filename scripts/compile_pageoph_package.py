import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def replace_frac(s):
    while r'\frac' in s:
        idx = s.find(r'\frac')
        open1 = s.find('{', idx)
        if open1 == -1:
            break
        cnt = 1
        pos = open1 + 1
        while pos < len(s) and cnt > 0:
            if s[pos] == '{': cnt += 1
            elif s[pos] == '}': cnt -= 1
            pos += 1
        if cnt != 0: break
        close1 = pos - 1
        num = s[open1+1:close1]
        
        open2 = s.find('{', close1)
        if open2 == -1: break
        cnt = 1
        pos = open2 + 1
        while pos < len(s) and cnt > 0:
            if s[pos] == '{': cnt += 1
            elif s[pos] == '}': cnt -= 1
            pos += 1
        if cnt != 0: break
        close2 = pos - 1
        den = s[open2+1:close2]
        
        s = s[:idx] + f"({num}) / ({den})" + s[close2+1:]
    return s

def clean_math_text(s):
    for _ in range(5):
        s = re.sub(r'\\(?:text|mathbf|mathrm|mathit)\{([^}]*)\}', r'\1', s)
        s = replace_frac(s)
        s = re.sub(r'\\bar\{([^{}]+)\}', r'\1̄', s)
    
    s = re.sub(r'\\left\(', '(', s)
    s = re.sub(r'\\right\)', ')', s)
    s = re.sub(r'\\left\[', '[', s)
    s = re.sub(r'\\right\]', ']', s)
    s = re.sub(r'\\left\\\{', '{', s)
    s = re.sub(r'\\right\\\}', '}', s)
    s = re.sub(r'\\\{', '{', s)
    s = re.sub(r'\\\}', '}', s)
    
    replacements = [
        (r'\\begin\{cases\}', '{\n  '),
        (r'\\end\{cases\}', '\n}'),
        (r'\\\\', '\n  '),
        (r'\\quad', '  '),
        (r'\\qquad', '    '),
        (r'\\implies', '⟹'),
        (r'\\min\b', 'min'),
        (r'\\max\b', 'max'),
        (r'\\sum\b', '∑'),
        (r'\\_', '_'),
        (r'\\Delta', 'Δ'),
        (r'\\sigma\^0', 'σ⁰'),
        (r'\\sigma', 'σ'),
        (r'\\theta', 'θ'),
        (r'\\pi', 'π'),
        (r'\\rho', 'ρ'),
        (r'\\times', '×'),
        (r'\\ge', '≥'),
        (r'\\le', '≤'),
        (r'\\approx', '≈'),
        (r'\\sim', '~'),
        (r'\\pm', '±'),
        (r'\\cdot', '·'),
        (r'\\circ', '°'),
        (r'\\land', '∧'),
        (r'\\lor', '∨'),
        (r'\\in', '∈'),
        (r'\\rightarrow', '→'),
        (r'\\to\b', '→'),
        (r'\\dagger', '†'),
        (r'\\dots', '...'),
        (r'--', '–'),
        (r'\\%', '%'),
    ]
    for pattern, repl in replacements:
        s = re.sub(pattern, repl, s)
    return s

def clean_plain_text(s):
    if '$' in s:
        s = clean_math_text(s)
    s = s.replace('$', '')
    for _ in range(5):
        s = re.sub(r'\\bar\{([^{}]+)\}', r'\1̄', s)
        s = re.sub(r'\\(?:text|mathbf|mathrm|mathit)\{([^}]*)\}', r'\1', s)
    replacements = [
        (r'\\text\{--\}', '–'),
        (r'\\text\{([^}]*)\}', r'\1'),
        (r'\\implies', '⟹'),
        (r'\\min\b', 'min'),
        (r'\\max\b', 'max'),
        (r'\\sum\b', '∑'),
        (r'\\ge', '≥'),
        (r'\\le', '≤'),
        (r'\\times', '×'),
        (r'\\Delta', 'Δ'),
        (r'\\sigma\^0', 'σ⁰'),
        (r'\\sigma', 'σ'),
        (r'\\theta', 'θ'),
        (r'\\pi', 'π'),
        (r'\\rho', 'ρ'),
        (r'\\pm', '±'),
        (r'\\approx', '≈'),
        (r'\\dagger', '†'),
        (r'\\rightarrow', '→'),
        (r'\\to\b', '→'),
        (r'\\%', '%'),
    ]
    for pattern, repl in replacements:
        s = re.sub(pattern, repl, s)
    return s

def add_inline_formatting(paragraph, text, base_font_size=10.5, is_bold_base=False, is_italic_base=False, font_name="Calibri", font_color=None):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$[^\$]+?\$)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = font_name
        run.font.size = Pt(base_font_size)
        if font_color:
            run.font.color.rgb = font_color

        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = clean_plain_text(token[2:-2])
            run.bold = True
            run.italic = is_italic_base
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = clean_plain_text(token[1:-1])
            run.bold = is_bold_base
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size - 1)
            run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        elif token.startswith('$') and token.endswith('$') and len(token) >= 2:
            run.text = clean_math_text(token[1:-1])
            run.font.name = "Cambria Math"
            run.italic = True
        else:
            run.text = clean_plain_text(token)
            run.bold = is_bold_base
            run.italic = is_italic_base

def markdown_to_docx(md_path, docx_path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lines = []
    in_table = False
    table_lines = []

    def flush_table(table_data):
        if not table_data:
            return
        rows = []
        for r_line in table_data:
            r_line = r_line.strip()
            if not r_line or set(r_line).issubset({'|', '-', ' ', ':'}):
                continue
            cells = [c.strip() for c in r_line.strip('|').split('|')]
            rows.append(cells)
        if not rows:
            return

        n_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=n_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        # Auto font size based on column count
        font_sz = 8.0 if n_cols >= 10 else (9.0 if n_cols >= 6 else 9.5)

        for i, row in enumerate(rows):
            for j in range(n_cols):
                cell_text = row[j] if j < len(row) else ""
                cell = table.cell(i, j)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                is_header = (i == 0)
                add_inline_formatting(p, cell_text, base_font_size=font_sz, is_bold_base=is_header)
                set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
                if is_header:
                    set_cell_background(cell, "1E3A8A") # Navy header
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                elif i % 2 == 1:
                    set_cell_background(cell, "F8FAFC") # Soft slate alternate

        doc.add_paragraph()

    for line in lines:
        raw_line = line.rstrip('\r\n')

        # Image check: ![caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', raw_line.strip())
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                doc.add_picture(img_path, width=Inches(6.4))
            continue

        # Code block handling
        if raw_line.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Inches(0.2)
                block_text = "\n".join(code_block_lines)
                run = p.add_run(block_text)
                run.font.name = "Consolas"
                run.font.size = Pt(8.0)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                code_block_lines = []
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
                code_block_lines = []
            continue

        if in_code_block:
            code_block_lines.append(raw_line)
            continue

        # Table row handling
        if '|' in raw_line and not raw_line.startswith('#'):
            in_table = True
            table_lines.append(raw_line)
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if not raw_line.strip():
            continue

        if raw_line.strip() in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 50)
            run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
            continue

        # Display math $$...$$
        if raw_line.startswith('$$') and raw_line.endswith('$$') and len(raw_line) >= 4:
            math_text = raw_line[2:-2].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.4)
            run = p.add_run(clean_math_text(math_text))
            run.font.name = "Cambria Math"
            run.font.size = Pt(11)
            run.italic = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            continue

        # Headings
        if raw_line.startswith('# '):
            text = raw_line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            add_inline_formatting(p, text, base_font_size=18, is_bold_base=True, font_color=RGBColor(0x0F, 0x17, 0x2A))
        elif raw_line.startswith('## '):
            text = raw_line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(13)
            p.paragraph_format.space_after = Pt(6)
            add_inline_formatting(p, text, base_font_size=14, is_bold_base=True, font_color=RGBColor(0x1E, 0x3A, 0x8A))
        elif raw_line.startswith('### '):
            text = raw_line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            add_inline_formatting(p, text, base_font_size=12, is_bold_base=True, font_color=RGBColor(0x03, 0x69, 0xA1))
        elif raw_line.startswith('#### '):
            text = raw_line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            add_inline_formatting(p, text, base_font_size=11, is_bold_base=True, font_color=RGBColor(0x33, 0x41, 0x55))
        elif raw_line.startswith('> '):
            text = raw_line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline_formatting(p, text, base_font_size=10, is_italic_base=True, font_color=RGBColor(0x33, 0x41, 0x55))
        elif raw_line.strip().startswith(('- ', '* ', '• ')):
            text = raw_line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_inline_formatting(p, text, base_font_size=10.5)
        elif re.match(r'^(\d+)[\.\)]\s+(.*)', raw_line.strip()):
            m = re.match(r'^(\d+)[\.\)]\s+(.*)', raw_line.strip())
            num_str = m.group(1)
            text = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run_num = p.add_run(f"{num_str}. ")
            run_num.font.name = "Calibri"
            run_num.font.size = Pt(10.5)
            run_num.bold = True
            add_inline_formatting(p, text, base_font_size=10.5)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_inline_formatting(p, raw_line, base_font_size=10.5)

    if in_table:
        flush_table(table_lines)

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

if __name__ == '__main__':
    md1 = '/Users/shubhamsharma/.gemini/antigravity/brain/e2a61b8e-ff9e-4a79-a6a7-b00c9c265b4f/REVISED_MANUSCRIPT_PAGEOPH.md'
    out1 = '/Users/shubhamsharma/.gemini/antigravity/brain/e2a61b8e-ff9e-4a79-a6a7-b00c9c265b4f/REVISED_MANUSCRIPT_PAGEOPH.docx'
    markdown_to_docx(md1, out1)

    md2 = '/Users/shubhamsharma/.gemini/antigravity/brain/e2a61b8e-ff9e-4a79-a6a7-b00c9c265b4f/RESPONSE_TO_REVIEWERS_PAGEOPH.md'
    out2 = '/Users/shubhamsharma/.gemini/antigravity/brain/e2a61b8e-ff9e-4a79-a6a7-b00c9c265b4f/RESPONSE_TO_REVIEWERS_PAGEOPH.docx'
    markdown_to_docx(md2, out2)
